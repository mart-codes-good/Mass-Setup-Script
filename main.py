"""
Lector welcome sheet generator -> Word (.docx)

Usage examples:
    python lector_sheet.py                          # next Sunday, asks for the rest
    python lector_sheet.py --date 2026-10-11 --time "12:30 PM" \
        --collection "The Needs of the Church in Canada" \
        --special "John Smith" --soul "Mary Brown" --soul "Peter Lee"
    python lector_sheet.py --title "Thanksgiving Day"   # manual override of the celebration name
    python lector_sheet.py --prayer michael             # add the St. Michael prayer

Anything not given on the command line is asked for interactively
(press Enter to accept the [default] shown).

Requires: pip install python-docx
Optional: litcal_roman (used for the celebration name if installed)
"""

import argparse
import datetime as dt
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, Inches

try:
    from litcal_roman import get_liturgical_day, CalendarConfig
except ImportError:
    get_liturgical_day = None
    CalendarConfig = None


# --------------------------------------------------------------------------
# Settings you'll rarely change
# --------------------------------------------------------------------------
PARISH_NAME = "St. Josephine Bakhita Parish"
DEFAULT_TIME = "12:30 PM"
FONT = "Cambria"

# Swappable prayer spot: add more prayers here and pick them with --prayer
PRAYERS = {
    "michael": """St. Michael the Archangel, defend us in battle.
Be our protection against the wickedness and snares of the Devil.
May God rebuke him, we humbly pray;
and do thou, O Prince of the Heavenly Host, by the power of God,
cast into Hell Satan and all the evil spirits,
who prowl about the world seeking the ruin of souls.
Amen""",
}


# --------------------------------------------------------------------------
# Liturgical name lookup
# --------------------------------------------------------------------------
_ONES = ["", "First", "Second", "Third", "Fourth", "Fifth", "Sixth", "Seventh",
         "Eighth", "Ninth", "Tenth", "Eleventh", "Twelfth", "Thirteenth",
         "Fourteenth", "Fifteenth", "Sixteenth", "Seventeenth", "Eighteenth",
         "Nineteenth"]
_TENS = {20: ("Twenty", "Twentieth"), 30: ("Thirty", "Thirtieth")}


def ordinal_words(n: int) -> str:
    """5 -> 'Fifth', 27 -> 'Twenty-Seventh' (1..39)."""
    if n < 20:
        return _ONES[n]
    tens, ones = divmod(n, 10)
    base, exact = _TENS[tens * 10]
    return exact if ones == 0 else f"{base}-{_ONES[ones]}"


def spell_out_ordinals(text: str) -> str:
    """'27th Sunday in Ordinary Time' -> 'Twenty-Seventh Sunday in Ordinary Time'."""
    return re.sub(r"\b(\d{1,2})(st|nd|rd|th)\b",
                  lambda m: ordinal_words(int(m.group(1))), text)


def easter(year: int) -> dt.date:
    """Gregorian Easter (anonymous algorithm)."""
    a, b, c = year % 19, year // 100, year % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month, day = divmod(h + l - 7 * m + 114, 31)
    return dt.date(year, month, day + 1)


def fallback_sunday_name(date: dt.date):
    """Built-in fallback for Sundays in Ordinary Time after Pentecost.
    Returns None when it can't tell (other seasons) -> you'll be asked."""
    if date.weekday() != 6:
        return None
    pentecost = easter(date.year) + dt.timedelta(days=49)
    christmas = dt.date(date.year, 12, 25)
    advent1 = christmas - dt.timedelta(days=christmas.weekday() + 1 + 21)
    christ_the_king = advent1 - dt.timedelta(days=7)
    if date == christ_the_king:
        return "Solemnity of Our Lord Jesus Christ, King of the Universe"
    # Trinity Sunday and Corpus Christi follow Pentecost; leave those to the library/user
    if pentecost + dt.timedelta(days=14) < date < christ_the_king:
        week = 34 - (christ_the_king - date).days // 7
        return f"{ordinal_words(week)} Sunday in Ordinary Time"
    return None


def lookup_celebration(date: dt.date):
    """Try litcal_roman first, then the built-in fallback."""
    if get_liturgical_day is not None:
        try:
            day = get_liturgical_day(dt.datetime.combine(date, dt.time()),
                                     CalendarConfig.us())
            for attr in ("name", "title", "celebration", "full_name"):
                value = getattr(day, attr, None)
                if value:
                    return spell_out_ordinals(str(value))
            return spell_out_ordinals(str(day))
        except Exception as exc:  # library bug / unsupported date -> fall back
            print(f"  (litcal_roman failed: {exc}; using fallback)")
    return fallback_sunday_name(date)


# --------------------------------------------------------------------------
# Word document
# --------------------------------------------------------------------------
def _para(doc_or_cell, text="", *, size=14, bold=False, italic=False,
          underline=False, highlight=False, align=WD_ALIGN_PARAGRAPH.CENTER,
          space_after=10):
    p = doc_or_cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        _run(p, text, size=size, bold=bold, italic=italic,
             underline=underline, highlight=highlight)
    return p


def _run(p, text, *, size=14, bold=False, italic=False, underline=False,
         highlight=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold, r.italic, r.underline = bold, italic, underline
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def _first_para(cell):
    """Reuse the empty paragraph every new cell starts with."""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    return p


def build_document(*, date, celebration, masses, collection, prayer=None):
    """masses: list of dicts {time, special: [..], souls: [..]}"""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.8)

    _para(doc, f"Welcome to {PARISH_NAME}", size=16, bold=True)
    _para(doc, f"On the {celebration}", size=15, underline=True, space_after=16)

    p = _para(doc, space_after=16)
    _run(p, "I, ")
    _run(p, "_" * 28)
    _run(p, ", am the Lector for this Mass.")

    _para(doc, "This Mass Celebration is for...")

    # --- Mass / intentions table ---
    table = doc.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(1.9), Inches(1.4), Inches(3.6))
    for col, w in zip(table.columns, widths):
        col.width = w
    day_label = f"{date:%A}\n({date:%B} {date.day})"

    for mass in masses:
        cells = table.add_row().cells
        for c, w in zip(cells, widths):
            c.width = w
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        p = _first_para(cells[0])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, day_label, size=15, bold=True)

        p = _first_para(cells[1])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, mass["time"], size=15, bold=True, highlight=True)

        cell = cells[2]
        p = _first_para(cell)
        first = True
        for name in mass.get("special", []):
            if not first:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
            _run(p, f"Special Intention for {name}")
            first = False
        if mass.get("souls"):
            if not first:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
            _run(p, "For the Souls of...")
            for name in mass["souls"]:
                q = cell.add_paragraph()
                q.paragraph_format.space_after = Pt(2)
                _run(q, f"+ {name}", size=13)
            first = False
        if first:  # no intentions given
            _run(p, "")
        cell.add_paragraph()  # breathing room at bottom of the cell

    _para(doc, space_after=6)

    # --- Second offering box (single bordered cell) ---
    if collection:
        box = doc.add_table(rows=1, cols=1)
        box.style = "Table Grid"
        box.alignment = WD_TABLE_ALIGNMENT.CENTER
        box.autofit = False
        box.columns[0].width = Inches(5.2)
        cell = box.rows[0].cells[0]
        cell.width = Inches(5.2)
        p = _first_para(cell)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, "There will be a Second Offering after Holy Communion for")
        _para(cell, collection, size=16, bold=True, highlight=True, space_after=2)
        _para(doc, space_after=6)

    _para(doc, "Let us pause for a few moments of silence as we prepare to "
               "celebrate the Holy Eucharist", size=13, space_after=24)
    _para(doc, "* Pause for a moment *", italic=True, space_after=24)
    _para(doc, "Please Stand for the Opening Hymn!")

    # --- Swappable prayer spot ---
    if prayer:
        _para(doc, space_after=12)
        for line in prayer.strip().splitlines():
            _para(doc, line, size=13, italic=True, space_after=0)

    return doc


# --------------------------------------------------------------------------
# Input handling (CLI args + interactive manual override)
# --------------------------------------------------------------------------
def ask(prompt, default=""):
    shown = f" [{default}]" if default else ""
    answer = input(f"{prompt}{shown}: ").strip()
    return answer or default


def ask_list(prompt):
    print(f"{prompt} (one per line, blank line to finish)")
    items = []
    while (line := input("  > ").strip()):
        items.append(line)
    return items


def next_sunday(today=None):
    today = today or dt.date.today()
    return today + dt.timedelta(days=(6 - today.weekday()) % 7 or 7)


def main():
    ap = argparse.ArgumentParser(description="Generate the lector welcome sheet.")
    ap.add_argument("--date", help="YYYY-MM-DD (default: next Sunday)")
    ap.add_argument("--title", help="Override celebration, e.g. 'Twenty-Eighth Sunday in Ordinary Time'")
    ap.add_argument("--time", help=f"Mass time (default {DEFAULT_TIME})")
    ap.add_argument("--collection", help="Second offering for...")
    ap.add_argument("--special", action="append", help="Special intention name (repeatable)")
    ap.add_argument("--soul", action="append", help="Name of deceased (repeatable)")
    ap.add_argument("--prayer", choices=sorted(PRAYERS), help="Add a prayer at the end")
    ap.add_argument("--out", help="Output file name")
    ap.add_argument("-y", "--no-prompt", action="store_true",
                    help="Don't ask anything; use args/defaults only")
    args = ap.parse_args()
    interactive = not args.no_prompt

    # Date
    date = dt.date.fromisoformat(args.date) if args.date else next_sunday()
    if interactive and not args.date:
        date = dt.date.fromisoformat(ask("Date (YYYY-MM-DD)", date.isoformat()))

    # Celebration name: --title wins, else lookup, and always confirmable
    celebration = args.title or lookup_celebration(date)
    if interactive and not args.title:
        celebration = ask("Celebration", celebration or "")
    if not celebration:
        raise SystemExit("No celebration name found; pass --title.")

    time = args.time or (ask("Mass time", DEFAULT_TIME) if interactive else DEFAULT_TIME)
    collection = args.collection if args.collection is not None else (
        ask("Second offering for (blank = none)") if interactive else "")
    special = args.special if args.special is not None else (
        ask_list("Special intentions") if interactive else [])
    souls = args.soul if args.soul is not None else (
        ask_list("For the souls of") if interactive else [])

    doc = build_document(
        date=date,
        celebration=celebration,
        masses=[{"time": time, "special": special, "souls": souls}],
        collection=collection,
        prayer=PRAYERS.get(args.prayer) if args.prayer else None,
    )
    out = args.out or f"Lector_Sheet_{date.isoformat()}.docx"
    doc.save(out)
    print(f"Saved {out}")


if __name__ == "__main__":
    main()